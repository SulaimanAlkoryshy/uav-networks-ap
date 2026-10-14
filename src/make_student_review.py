"""
Build student_review.docx — a mid-level explanation of the AP project, organised Stage -> steps.

Language: Arabic text with the ML/AI terms kept in English. Every paragraph is set right-to-left
(bidi) with an Arabic-capable font, so mixed Arabic/English reads correctly in Word.
The numbers are read from data/processed/*.csv, so the document matches the latest run.

Run:
    python src/make_student_review.py          # writes student_review.docx
Then to PDF (Windows + Word):
    powershell -File src/docx_to_pdf.ps1 student_review.docx
"""

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Arial"
BLUE = RGBColor(0x1F, 0x3A, 0x6E)
D = "data/processed"


# ==========================================================================
# STEP 1 — make_rtl(): turn one paragraph into a right-to-left Arabic paragraph
# ==========================================================================
def make_rtl(p, size, bold=False, color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))            # paragraph direction = right-to-left
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "start")                 # "start" of an RTL paragraph = the right side
    pPr.append(jc)
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rPr.append(OxmlElement("w:rtl"))
        fonts = rPr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rPr.append(fonts)
        fonts.set(qn("w:cs"), FONT)              # font for Arabic (complex script)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:bidi"), "ar-SA")
        rPr.append(lang)


# ==========================================================================
# STEP 2 — small writers: heading / paragraph / bullet
# ==========================================================================
def heading(doc, text, size=15):
    p = doc.add_paragraph()
    p.add_run(text)
    make_rtl(p, size, bold=True, color=BLUE)


def para(doc, text, size=11.5, bold=False):
    p = doc.add_paragraph()
    p.add_run(text)
    make_rtl(p, size, bold=bold)


def bullet(doc, text):
    p = doc.add_paragraph()
    p.add_run("•  " + text)
    make_rtl(p, 11.5)


# ==========================================================================
# STEP 3 — read the numbers from the saved results
# ==========================================================================
def load_numbers():
    n = {}
    s2 = pd.read_csv(f"{D}/stage2_results.csv").set_index(["path", "strategy"])["macro_f1"]
    n["s2"] = s2
    s3 = pd.read_csv(f"{D}/stage3_binary_results.csv").set_index("path")
    n["bin_search"] = s3.loc["search", "cv_macroF1"]
    n["bin_base"] = s3.loc["search", "baseline_allAttack_macroF1"]
    cap = pd.read_csv(f"{D}/stage3_capture_check.csv").set_index("columns_dropped")["test_macroF1"]
    n["cap6"] = cap.loc[6]
    s4 = pd.read_csv(f"{D}/stage4_multiclass_results.csv").set_index("path")["cv_macroF1"]
    n["mc_report"], n["mc_search"] = s4["report"], s4["search"]
    n["pc"] = pd.read_csv(f"{D}/stage4_per_class_f1.csv").set_index("class")
    s5 = pd.read_csv(f"{D}/stage5_benchmark.csv")
    n["s5"] = s5.pivot(index="model", columns="path", values="test_macroF1")
    n["leak"] = pd.read_csv(f"{D}/stage1_shortcut_check.csv").iloc[0]
    n["cost"] = pd.read_csv(f"{D}/stage6_cost.csv").set_index("path")
    n["perm_report"] = list(pd.read_csv(f"{D}/stage6_report_permutation.csv")["feature"].head(3))
    n["perm_search"] = list(pd.read_csv(f"{D}/stage6_search_permutation.csv")["feature"].head(3))
    return n


# ==========================================================================
# STEP 4 — write the document, Stage -> steps
# ==========================================================================
def main():
    n = load_numbers()
    pc, s5 = n["pc"], n["s5"]
    best = s5["search"].idxmax()
    doc = Document()

    heading(doc, "شرح المشروع — كشف التسلل في شبكات الطائرات بدون طيار (سيناريو Access Point)", 17)
    para(doc, "المشروع نظام كشف تسلل (Intrusion Detection System — IDS) يستخدم Classical Machine Learning "
              "على داتا UAV-NIDD، وتحديدًا السيناريو الثاني (Case 2): الـ Access Point هو النقطة المخترقة. "
              "الداتا عبارة عن حزم Wi-Fi (packet-level): 631,356 حزمة، و10 أنواع هجوم + Normal.")
    para(doc, "نمشي بمسارين في كل المراحل: مسار report (نطبّق اللي يقوله التقرير حرفيًا)، ومسار search "
              "(النسخة الصادقة اللي تحذف الأعمدة اللي تكشف الجواب بالغش). الهدف رقم نقدر نثق فيه، "
              "مو أكبر رقم.")

    # ---- Stage 1 ----
    heading(doc, "المرحلة 1 — تجهيز البيانات واختيار الميزات (Preprocessing & Feature Selection)")
    bullet(doc, "الخطوة 1: تحميل الملف. الملف Excel كبير، فحوّلناه مرة وحدة إلى CSV وصرنا نقرأ الـ CSV.")
    bullet(doc, "الخطوة 2: إصلاح الـ label. عمود الفئة اسمه Normal بالغلط فسمّيناه Class، وصحّحنا "
                "الخطأ الإملائي Ewil Twin إلى Evil Twin.")
    bullet(doc, "الخطوة 3: تحويل الأعمدة لأرقام. أسماء الأعمدة مزاحة عن محتواها، وفيه خلايا فيها IP أو hex. "
                "حوّلنا الـ hex لأرقام وحطينا -1 مكان أي نص. ما حذفنا الصفوف، لأن حذفها كان بيمسح "
                "100% من Evil Twin وMITM وReconnaissance وReplay.")
    bullet(doc, f"الخطوة 4: فحص الغش (Shortcut check). جرّبنا كل عمود لحاله: العمود "
                f"{n['leak']['column']} (وهو فعليًا وقت الالتقاط) لحاله أعطى macro-F1 = "
                f"{n['leak']['macro_f1_alone']:.2f}. السبب أن كل هجوم انسجل في وقت مختلف، فالوقت "
                "يكشف نوع الهجوم. هذا اسمه Data Leakage، فحذفنا أعمدة الوقت والترتيب وعناوين IP "
                "في مسار search.")
    bullet(doc, "الخطوة 5: تقسيم البيانات (Train/Test Split) 80/20 مع stratify وseed = 42.")
    bullet(doc, "الخطوة 6: حذف الميزات المتشابهة بـ PCC (|corr| > 0.85) على بيانات التدريب فقط.")
    bullet(doc, "الخطوة 7: اختيار أفضل 25 ميزة: مسار report بـ ANOVA، ومسار search بأهمية Random Forest. "
                "لاحظنا إن ANOVA اختار أعمدة الوقت والـ IP، يعني وقع في الـ leakage.")

    # ---- Stage 2 ----
    heading(doc, "المرحلة 2 — موازنة الفئات (Class Balancing)")
    bullet(doc, "المشكلة: Brute-Force فيها 172 ألف حزمة وReconnaissance فيها 40 بس (تقريبًا 4,300 : 1).")
    bullet(doc, "قارنّا ثلاث طرق: بدون موازنة (none)، وRandomUnderSampler ثم Borderline-SMOTE "
                "(under_smote)، والأوزان (class_weight).")
    bullet(doc, f"النتيجة في المسار الصادق: none = {n['s2'][('search', 'none')]:.3f}، "
                f"under_smote = {n['s2'][('search', 'under_smote')]:.3f}، "
                f"class_weight = {n['s2'][('search', 'class_weight')]:.3f}. يعني الفرق صغير جدًا، "
                "والزيادة +0.18 اللي يذكرها التقرير ما تكررت.")
    bullet(doc, "اخترنا under_smote لأنها ترفع الـ recall للهجمات الضعيفة (DDoS وMITM وReplay) "
                "وتصغّر بيانات التدريب.")

    # ---- Stage 3 ----
    heading(doc, "المرحلة 3 — تصنيف ثنائي (Binary: Normal vs Attack)")
    bullet(doc, "النموذج Random Forest، والتقييم 5-fold Cross-Validation، مع مقارنة بنموذج غبي "
                f"يقول دائمًا Attack (macro-F1 = {n['bin_base']:.2f}).")
    bullet(doc, f"النتيجة {n['bin_search']:.3f} حتى في المسار الصادق، وهذي إشارة خطر.")
    bullet(doc, "بحثنا عن السبب: كل حزم Normal قيمة frame.encap_type فيها 23، وما فيه ولا هجوم بهذي القيمة. "
                f"وحتى بعد حذف أقوى 6 أعمدة بقي الـ macro-F1 = {n['cap6']:.2f}. يعني Normal "
                "انسجلت بطريقة التقاط مختلفة، والنموذج يتعرّف على طريقة التسجيل مو على السلوك. "
                "فكتبنا النتيجة لكن ما اعتبرناها قدرة كشف حقيقية.")

    # ---- Stage 4 ----
    heading(doc, "المرحلة 4 — تصنيف متعدد (Multi-class: 11 فئة)")
    bullet(doc, f"5-fold CV macro-F1: مسار report = {n['mc_report']:.3f}، ومسار search = "
                f"{n['mc_search']:.3f}. الفرق جاي من أعمدة الوقت.")
    bullet(doc, "فئات ممتازة (F1 ≥ 0.99): Brute-Force وDoS وNormal وScanning وFake Landing.")
    bullet(doc, f"فئات ضعيفة: DDoS = {pc.loc['DDoS', 'search_f1']:.2f}، Replay = "
                f"{pc.loc['Replay', 'search_f1']:.2f}، MITM = {pc.loc['MITM', 'search_f1']:.2f}. "
                "DDoS تختلط مع DoS، وReplay وMITM شكلها قريب من الحزم العادية.")
    bullet(doc, "Reconnaissance نتيجتها 1.00 لكن من 32 صف بس للتدريب، فما نعتمد عليها.")

    # ---- Stage 5 ----
    heading(doc, "المرحلة 5 — مقارنة النماذج (Model Benchmarking)")
    for m in s5.sort_values("search", ascending=False).index:
        bullet(doc, f"{m}: report = {s5.loc[m, 'report']:.3f}، search = {s5.loc[m, 'search']:.3f}")
    bullet(doc, f"أفضل نموذج صادق هو {best} ({s5.loc[best, 'search']:.3f}). نماذج الأشجار أقوى بكثير "
                "من النماذج الخطية (Logistic Regression وLinear SVM). هدف 0.90 تحقق فقط مع الـ leakage.")

    # ---- Stage 6 ----
    heading(doc, "المرحلة 6 — التفسير والتكلفة وإعادة الإنتاج")
    bullet(doc, "أهم الميزات (Permutation Importance) في مسار report: " + "، ".join(n["perm_report"]) +
                ". كلها أعمدة وقت، وهذا دليل على الغش.")
    bullet(doc, "أهم الميزات في مسار search: " + "، ".join(n["perm_search"]) + ".")
    c = n["cost"].loc["search"]
    bullet(doc, f"التكلفة على CPU: التدريب {c['train (s)']} ثانية، والسرعة تقريبًا {c['rows/sec']} حزمة بالثانية.")
    bullet(doc, "إعادة الإنتاج (Reproducibility): نفس الـ seed يعطي نفس التوقعات بالضبط، و"
                "run_all.py يعيد كل الأرقام بأمر واحد.")

    # ---- conclusion ----
    heading(doc, "الخلاصة")
    para(doc, f"الأرقام العالية (~{n['mc_report']:.2f}) تعتمد على وقت الالتقاط. الرقم الصادق اللي نقدر نثق فيه "
              f"حوالي {n['mc_search']:.2f}–{s5.loc[best, 'search']:.2f} macro-F1، ونقاط الضعف DDoS وReplay وMITM. "
              "والخطوة الجاية المقترحة: تقسيم البيانات حسب الوقت (time-aware split) كاختبار أقوى ضد الـ leakage.")

    doc.save("student_review.docx")
    print("wrote student_review.docx")


if __name__ == "__main__":
    main()
